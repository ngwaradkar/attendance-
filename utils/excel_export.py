from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from utils.helpers import now

def excel_bytes(df,title,institution='Vidya Prabodhini',year='2026–27',date_range='All available records',threshold=75):
    wb=Workbook(); ws=wb.active; ws.title='Report'
    frame=df.drop(columns=[c for c in ['class_id','student_id','division_id','stream_id','year_id','session_id'] if c in df],errors='ignore')
    n=max(1,len(frame.columns))
    for row,text in [(1,institution),(2,title),(3,f'Academic year: {year} | {date_range}'),(4,f'Generated: {now()}')]:
        ws.cell(row,1,text); ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=n)
        ws.cell(row,1).font=Font(name='Calibri',size=18 if row==1 else 11,bold=row<3,color='17365D')
    for j,col in enumerate(frame.columns,1):
        cell=ws.cell(6,j,str(col)); cell.fill=PatternFill('solid',fgColor='17365D'); cell.font=Font(bold=True,color='FFFFFF'); cell.alignment=Alignment(wrap_text=True)
    thin=Side(style='thin',color='E2E8F0')
    for i,row in enumerate(frame.itertuples(index=False,name=None),7):
        for j,value in enumerate(row,1):
            if hasattr(value,'item'): value=value.item()
            if value is None or str(value) in ['nan','<NA>','NaT']: value=''
            # Prevent uploaded text from becoming executable spreadsheet formulas.
            if isinstance(value,str) and value.startswith(('=','+','-','@')) and value!='-': value="'"+value
            cell=ws.cell(i,j,value); cell.font=Font(name='Calibri',size=11); cell.border=Border(bottom=thin)
            if i%2: cell.fill=PatternFill('solid',fgColor='F1F5FB')
            if '%' in str(frame.columns[j-1]): cell.number_format='0.00"%"'
    for j,col in enumerate(frame.columns,1):
        width=max([len(str(col))]+[len(str(v)) for v in frame.iloc[:200,j-1]])+3
        ws.column_dimensions[get_column_letter(j)].width=min(42,max(11,width))
        if '%' in str(col) and len(frame):
            cells=f'{get_column_letter(j)}7:{get_column_letter(j)}{6+len(frame)}'
            ws.conditional_formatting.add(cells,CellIsRule(operator='lessThan',formula=[str(threshold)],fill=PatternFill('solid',fgColor='FEE2E2')))
    ws.freeze_panes='C7'; ws.auto_filter.ref=f'A6:{get_column_letter(n)}{max(6,6+len(frame))}'
    ws.sheet_view.showGridLines=False; ws.print_title_rows='1:6'; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A4; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    output=BytesIO(); wb.save(output); return output.getvalue()
